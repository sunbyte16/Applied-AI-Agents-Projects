import os
from pathlib import Path
from typing import List, Optional
from pydantic import BaseModel
import pypdf
import docx

from backend.ingestion.text_cleaner import TextCleaner
from backend.config import settings


class PageContent(BaseModel):
    text: str
    page_number: Optional[int] = None
    char_count: int


class LoadedDocument(BaseModel):
    document_name: str
    file_type: str
    pages: List[PageContent]
    total_characters: int


class DocumentLoader:
    """Loads and extracts text from various document formats (PDF, TXT, DOCX, MD)."""

    SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx", ".md"}

    @classmethod
    def validate_file(cls, file_path: Path) -> None:
        """Validates file existence, format, and size."""
        if not file_path.exists() or not file_path.is_file():
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = file_path.suffix.lower()
        if ext not in cls.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file format '{ext}'. Supported formats: {', '.join(sorted(cls.SUPPORTED_EXTENSIONS))}"
            )

        file_size_mb = file_path.stat().st_size / (1024 * 1024)
        if file_size_mb > settings.MAX_FILE_SIZE_MB:
            raise ValueError(
                f"File size ({file_size_mb:.2f} MB) exceeds maximum allowed limit ({settings.MAX_FILE_SIZE_MB} MB)."
            )

    @classmethod
    def load(cls, file_path: Path) -> LoadedDocument:
        """Extracts and cleans text from the given document file.

        Args:
            file_path: Path to the target document.

        Returns:
            LoadedDocument containing cleaned text broken down by page or section.
        """
        cls.validate_file(file_path)
        ext = file_path.suffix.lower()

        if ext == ".pdf":
            pages = cls._load_pdf(file_path)
        elif ext == ".docx":
            pages = cls._load_docx(file_path)
        elif ext in {".txt", ".md"}:
            pages = cls._load_text(file_path)
        else:
            raise ValueError(f"Unsupported extension: {ext}")

        total_chars = sum(p.char_count for p in pages)
        if total_chars == 0:
            raise ValueError(
                f"The document '{file_path.name}' contains no readable or extractable text."
            )

        return LoadedDocument(
            document_name=file_path.name,
            file_type=ext.lstrip("."),
            pages=pages,
            total_characters=total_chars
        )

    @classmethod
    def _load_pdf(cls, file_path: Path) -> List[PageContent]:
        pages = []
        try:
            reader = pypdf.PdfReader(str(file_path))
            for idx, page in enumerate(reader.pages):
                raw_text = page.extract_text() or ""
                cleaned = TextCleaner.clean(raw_text)
                if cleaned:
                    pages.append(
                        PageContent(
                            text=cleaned,
                            page_number=idx + 1,
                            char_count=len(cleaned)
                        )
                    )
        except Exception as e:
            if isinstance(e, ValueError):
                raise
            raise ValueError(f"Failed to parse PDF document: {str(e)}") from e

        return pages

    @classmethod
    def _load_docx(cls, file_path: Path) -> List[PageContent]:
        try:
            doc = docx.Document(str(file_path))
            paragraphs = []
            for p in doc.paragraphs:
                if p.text.strip():
                    paragraphs.append(p.text.strip())

            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            combined_text = "\n\n".join(paragraphs)
            cleaned = TextCleaner.clean(combined_text)
            if not cleaned:
                return []
            return [PageContent(text=cleaned, page_number=None, char_count=len(cleaned))]
        except Exception as e:
            if isinstance(e, ValueError):
                raise
            raise ValueError(f"Failed to parse DOCX document: {str(e)}") from e

    @classmethod
    def _load_text(cls, file_path: Path) -> List[PageContent]:
        raw_text = ""
        # Try UTF-8 first, then fallback to common encodings
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    raw_text = f.read()
                break
            except UnicodeDecodeError:
                continue

        cleaned = TextCleaner.clean(raw_text)
        if not cleaned:
            return []
        return [PageContent(text=cleaned, page_number=None, char_count=len(cleaned))]
